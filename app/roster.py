"""Read roster workbooks without executing formulas or changing source files."""
import base64
from collections import OrderedDict
from io import BytesIO
import re
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook


WEEKDAYS = '一二三四五六日'
PERIODS = '一二三四五'
TIMES = [('08:15','09:55'),('10:15','11:55'),('13:50','16:25'),('16:45','18:25'),('19:20','21:00')]


def slot_name(day, period):
    return f'周{WEEKDAYS[day-1]}-{PERIODS[period-1]}时段'


def category_slot(value):
    value = re.sub(r'\s+', '', str(value or ''))
    match = re.fullmatch(r'([一二三四五六七1-7])分队([一二三四五1-5])时段', value)
    if match:
        d,p = match.groups()
        return ('一二三四五六七'.index(d)+1 if not d.isdigit() else int(d),
                PERIODS.index(p)+1 if not p.isdigit() else int(p))
    match = re.fullmatch(r'周([一二三四五六日天])-([一二三四五])时段', value)
    if match:
        d,p=match.groups()
        return (WEEKDAYS.index('日' if d=='天' else d)+1, PERIODS.index(p)+1)
    return None


def parse_workbook(encoded, sheet_name=None, map_categories=True):
    try:
        raw=base64.b64decode(encoded, validate=True)
        if len(raw)>8*1024*1024:
            raise ValueError('文件不能超过 8 MB')
        with ZipFile(BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist())>40*1024*1024 or len(archive.infolist())>2000:
                raise ValueError('工作簿内容过大，请精简后导入')
        book=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ValueError('请上传有效的 .xlsx 文件（不超过 8 MB，解压内容不超过 40 MB）') from exc
    try:
        sheets=book.sheetnames
        chosen=sheet_name or sheets[0]
        if chosen not in sheets:
            raise ValueError('所选工作表不存在')
        sheet=book[chosen]
        if sheet.max_row is None or sheet.max_column is None:
            sheet.calculate_dimension(force=True)
        if sheet.max_row>10000 or sheet.max_column>100:
            raise ValueError('工作表不能超过 10000 行、100 列')
        header=None
        people=OrderedDict()
        errors=[]
        warnings=[]
        if '退队' in chosen:
            warnings.append({'row':0,'student_id':'','message':'当前选择的是退队名单；本导入会建立在队账号，不会导入退队状态，请确认是否选错工作表'})
        source_rows=0
        for number,row in enumerate(sheet.iter_rows(),1):
            if number>10000:
                raise ValueError('工作表行数过多')
            if header is None:
                labels={str(c.value).strip():i for i,c in enumerate(row) if c.value is not None}
                if '学号' in labels and '姓名' in labels:
                    header=labels
                elif number>=30:
                    raise ValueError('前 30 行未找到“学号”和“姓名”表头')
                continue
            def value(key):
                cell=row[header[key]] if key in header else None
                if cell is None or cell.value is None:
                    return ''
                if cell.data_type=='f':
                    raise ValueError(f'第 {number} 行的{key}是公式，请改为明确的文本值')
                v=cell.value
                if key=='学号' and isinstance(v,(int,float)):
                    if v!=int(v) or abs(v)>=10**15:
                        raise ValueError(f'第 {number} 行学号可能存在精度损失，请改用文本学号')
                    return str(int(v))
                return str(v).strip()
            try:
                sid,name=value('学号'),value('姓名')
                if not sid and not name:
                    continue
                source_rows+=1
                if not sid:
                    warnings.append({'row':number,'student_id':'','message':f'第 {number} 行缺少学号，已跳过，不建立账号'})
                    continue
                if not re.fullmatch(r'[A-Za-z0-9_-]{1,32}',sid) or not name or len(name)>100:
                    raise ValueError(f'第 {number} 行：请填写有效学号和姓名')
                college=value('学院')
                if len(college)>100:
                    raise ValueError(f'第 {number} 行学院超过 100 字')
                category=value('分类') or value('时段')
                pair=category_slot(category) if map_categories else None
                if map_categories and not pair:
                    warnings.append({'row':number,'student_id':sid,'message':f'分类“{category or "空白"}”未对应固定时段，将只导入账号，不增加该行时段归属'})
                person=people.setdefault(sid,{'student_id':sid,'name':name,'college':college,'slots':[],'rows':[]})
                if person['name']!=name:
                    raise ValueError(f'第 {number} 行：学号 {sid} 与第 {person["rows"][0]} 行的姓名不同，请先核对')
                if college and person['college'] and college!=person['college']:
                    warnings.append({'row':number,'student_id':sid,'message':'同一学号的学院不一致，采用首次非空学院'})
                person['college']=person['college'] or college
                person['rows'].append(number)
                if pair and list(pair) not in person['slots']:
                    person['slots'].append(list(pair))
            except ValueError as exc:
                errors.append({'row':number,'message':str(exc)})
        if header is None:
            raise ValueError('未找到“学号”和“姓名”表头')
        return {'sheets':sheets,'sheet':chosen,'source_rows':source_rows,'unique_students':len(people),
                'memberships':sum(len(p['slots']) for p in people.values()),'students':list(people.values()),
                'warnings':warnings,'errors':errors}
    finally:
        book.close()

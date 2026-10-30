"""Semester hours export; uses the application's installed XLSX dependency."""
from collections import defaultdict
from datetime import timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .db import one, all_rows


def export_hours(service, actor, semester_id, business_id=None, group_id=None):
    from .service import require
    with service.db.transaction() as c:
        viewer=service.actor(c, actor, ('SUPER','LEAD'))
        term = one(c, 'SELECT * FROM semester WHERE semester_id=%s', (semester_id,))
        require(term, '学期不存在', 404)
        students = all_rows(c, 'SELECT student_id,name,college,status FROM student ORDER BY student_id')
        from .business_schema import installed
        multi=installed(c)
        slots = all_rows(c, 'SELECT t.*,b.name AS business_name FROM service_slot t JOIN business b ON b.business_id=t.business_id' if multi else 'SELECT slot_id,weekday,period_no,team_name FROM service_slot')
        memberships = all_rows(c, 'SELECT student_id,slot_id FROM slot_member')
        schedules = all_rows(c, '''SELECT DISTINCT student_id,slot_id,'''+('group_id' if multi else 'NULL AS group_id')+''' FROM schedule
            WHERE schedule_date BETWEEN %s AND %s''', (term['start_date'], term['end_date']))
        records = all_rows(c, '''SELECT s.student_id,s.slot_id,s.schedule_date,r.service_minutes,'''+('s.group_id' if multi else 'NULL AS group_id')+'''
            FROM service_record r JOIN attendance a ON a.attendance_id=r.attendance_id
            JOIN schedule s ON s.schedule_id=a.schedule_id
            WHERE r.record_status='VALID' AND s.schedule_date BETWEEN %s AND %s
            ORDER BY s.schedule_id,r.record_id''', (term['start_date'], term['end_date']))
        group_names={r['group_id']:r['name'] for r in all_rows(c,'SELECT * FROM business_group')} if multi else {}
        if viewer['_role']=='LEAD':
            allowed=set(viewer['_lead_slots'])
            slots=[r for r in slots if r['slot_id'] in allowed]
            memberships=[r for r in memberships if r['slot_id'] in allowed]
            schedules=[r for r in schedules if r['slot_id'] in allowed]
            records=[r for r in records if r['slot_id'] in allowed]
            people={r['student_id'] for r in list(memberships)+list(schedules)+list(records)}
            students=[r for r in students if r['student_id'] in people]
    if group_id is not None:
        records=[r for r in records if (r.get('group_id') or 0)==group_id]
        schedules=[r for r in schedules if (r.get('group_id') or 0)==group_id]
        memberships=[]
        people={r['student_id'] for r in records+schedules};students=[s for s in students if s['student_id'] in people]
    virtual={}
    lookup={s['slot_id']:s for s in slots}
    for r in list(schedules)+list(records):
        gid=r.get('group_id')
        if gid:
            original=r['slot_id'];key=f'{original}:group:{gid}';r['slot_id']=key
            virtual[key]={**lookup[original],'slot_id':key,'group_id':gid,'group_name':group_names.get(gid,'未知小组')}
    slots=list(slots)+list(virtual.values())
    if business_id is not None:
        slots=[s for s in slots if s.get('business_id')==business_id]
        ids={s['slot_id'] for s in slots}
        memberships=[m for m in memberships if m['slot_id'] in ids]
        schedules=[m for m in schedules if m['slot_id'] in ids]
        records=[r for r in records if r['slot_id'] in ids]
        people={m['student_id'] for m in list(memberships)+list(schedules)+list(records)}
        students=[s for s in students if s['student_id'] in people]
    return build_workbook(term, students, slots, list(memberships) + list(schedules), records, service.clock())


def build_workbook(term, students, slots, memberships, records, now):
    """Group minutes before converting to hours; never duplicate multi-slot hours."""
    anchor = term['start_date'] - timedelta(days=term['start_date'].weekday())
    weeks = term['end_week']
    lookup = {r['slot_id']: r for r in slots}
    belongs = defaultdict(set)
    totals = defaultdict(int)
    team_totals = defaultdict(int)
    multi=any('business_name' in s for s in slots)
    def group_key(slot):
        day=slot.get('weekday') or 0
        if not multi:return day
        name=slot.get('business_name','咨询服务')
        group=(f'第{"一二三四五六七"[day-1]}分队' if day else slot.get('team_name','其他班次') if slot.get('schedule_kind')=='WEEKLY' else slot.get('group_name') or '未分组' if slot.get('schedule_kind')=='EVENT' else '其他班次')
        return (name,group)
    for r in memberships:
        belongs[r['student_id']].add(r['slot_id'])
    for r in records:
        week = (r['schedule_date'] - anchor).days // 7
        if not 0 <= week < weeks or not term['start_date'] <= r['schedule_date'] <= term['end_date']:
            continue
        belongs[r['student_id']].add(r['slot_id'])
        totals[r['student_id'], r['slot_id'], week] += int(r['service_minutes'])
        team_totals[group_key(lookup.get(r['slot_id'], {})), week] += int(r['service_minutes'])
    headers = [f'第 {i+1} 周\n{anchor+timedelta(weeks=i):%Y-%m-%d}～{anchor+timedelta(weeks=i,days=6):%Y-%m-%d}' for i in range(weeks)]
    def hours(value, week):
        return value / 60 if value or anchor + timedelta(weeks=week) <= now.date() else None
    def category(slot):
        if slot is None: return '未分配时段'
        s = lookup.get(slot, {})
        return (f'{"一二三四五六七"[s["weekday"]-1]}分队{"一二三四五"[s["period_no"]-1] if 1<=s["period_no"]<=5 else str(s["period_no"])}时段'
                if s.get('weekday') else (s.get('group_name') or '未分组')+' · '+s.get('team_name',f'班次 {slot}') if s.get('schedule_kind')=='EVENT' else s.get('team_name', f'班次 {slot}'))
    def sort_slot(slot):
        s = lookup.get(slot, {})
        return (s.get('weekday') or 99, s.get('period_no') or 99, str(slot or 0))
    book = Workbook()
    book.remove(book.active)
    note = (f'单位：小时；导出时间：{now:%Y-%m-%d %H:%M}。仅计有效服务记录，含人工修正；'
            '当前周为已认定时长，未来周留空。第1周按开工日所在周一至周日，开工前不计时。'
            '开心花园以周日为统计归属日，并非实际服务日。'
            '一人多时段分行，个人学期总计为本行时段合计，跨行相加即个人总时长；分队汇总含退队人员历史时长。')
    def sheet(name, labels):
        s = book.create_sheet(name)
        last = len(labels)
        s.merge_cells(start_row=1,start_column=1,end_row=2,end_column=last)
        s.cell(1,1,f'{term["name"]} · {name}（{term["start_date"]:%Y-%m-%d} 至 {term["end_date"]:%Y-%m-%d}）')
        s.merge_cells(start_row=3,start_column=1,end_row=3,end_column=last)
        s.cell(3,1,note)
        for col,label in enumerate(labels,1):
            s.merge_cells(start_row=4,start_column=col,end_row=6,end_column=col)
            s.cell(4,col,label)
        s.freeze_panes='F7' if last==weeks+7 else 'B7'
        s.sheet_view.showGridLines=False
        s.row_dimensions[1].height=25
        s.row_dimensions[3].height=44
        for n in (4,5,6): s.row_dimensions[n].height=22
        return s
    labels=['姓名','业务','分类','学院','学号']+headers+['个人学期总计','个人备注']
    main=sheet('志愿时长统计',labels)
    summary=sheet('分队时长统计',['分队']+headers+['分队总计'])
    quit_sheet=sheet('退队统计',labels)
    def append(s, values):
        row=max(7,s.max_row+1)
        for col,value in enumerate(values,1):s.cell(row,col,value)
    entries=[(s,slot) for s in students for slot in (belongs[s['student_id']] or {None})]
    entries.sort(key=lambda item:(sort_slot(item[1]),item[0]['student_id']))
    for student,slot in entries:
        sid=student['student_id']
        s=quit_sheet if student['status']=='QUIT' else main
        values=[hours(totals[sid,slot,i],i) for i in range(weeks)]
        append(s,[student['name'],lookup.get(slot,{}).get('business_name','咨询服务'),category(slot),student.get('college') or '',str(sid)]+values+
                 [sum(totals[sid,slot,i] for i in range(weeks))/60,'已退队' if student['status']=='QUIT' else '已停用' if student['status']=='DISABLED' else ''])
    groups=sorted({key[0] for key in team_totals}) if multi else list(range(1,8))+([0] if any(team_totals[0,i] for i in range(weeks)) else [])
    for group in groups:
        name=' · '.join(group) if multi else (f'第{"一二三四五六七"[group-1]}分队' if group else '其他班次')
        append(summary,[name]+[hours(team_totals[group,i],i) for i in range(weeks)]+[sum(team_totals[group,i] for i in range(weeks))/60])
    append(summary,['全队合计']+[hours(sum(team_totals[g,i] for g in groups),i) for i in range(weeks)]+[sum(team_totals.values())/60])
    for s in book:
        width=s.max_column
        s.print_title_rows='1:6'
        s.sheet_properties.pageSetUpPr.fitToPage=True
        s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3
        s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
        for col in range(1,width+1):
            s.column_dimensions[get_column_letter(col)].width=25 if s==summary or col>5 else [18,14,23,30,23][col-1]
        for row in s:
            for cell in row:
                if cell.__class__.__name__=='MergedCell': continue
                # User-entered strings must remain literal, even if they begin with '='.
                if isinstance(cell.value,str): cell.data_type='s'
                cell.font=Font(name='微软雅黑',size=11,color='203D39',bold=cell.row<7)
                cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                if cell.row<7:
                    cell.fill=PatternFill('solid',fgColor='DFEDE4' if cell.row>=4 else 'F1F6F3')
                else:
                    cell.border=Border(bottom=Side(style='hair',color='D2DDD6'))
                    if cell.row%2:cell.fill=PatternFill('solid',fgColor='F5F8F6')
                    if isinstance(cell.value,(int,float)):cell.number_format='0.00'
            if row[0].row>=7:s.row_dimensions[row[0].row].height=30
        if s!=summary:
            for row in range(7,s.max_row+1):s.cell(row,5).number_format='0'
    output=BytesIO();book.save(output)
    return output.getvalue()

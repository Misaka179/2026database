"""Identity-safe business/group roster import; never recognizes service hours."""
import base64,hashlib,json
from io import BytesIO
from zipfile import ZipFile
from openpyxl import load_workbook
from .db import one,all_rows
from .security import hash_password

SHEET_GROUPS={'综运':'综合运营中心','图志之家':'图志之家','活策':'活动策划中心','物资采购':'物资采购中心','宣传':'宣传推广中心','音视频':'音视频制作中心','摄影摄像':'摄影摄像中心','青广':'青广干事'}

def parse_roster(encoded):
    from .service import require,DomainError
    try:
        raw=base64.b64decode(encoded,validate=True);require(len(raw)<=8*1024*1024,'文件不能超过8MB',422)
        with ZipFile(BytesIO(raw)) as z:require(len(z.infolist())<=2000 and sum(f.file_size for f in z.infolist())<=40*1024*1024,'工作簿解压内容过大',422)
        book=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        sheets=[];rows=[]
        for sheet in book:
            headers=None;count=0
            for number,cells in enumerate(sheet.iter_rows(),1):
                values=[str(c.value).strip() if c.value is not None else '' for c in cells]
                if '姓名' in values and '学号' in values:headers={v:i for i,v in enumerate(values) if v};continue
                if headers is None or not any(values):continue
                name=values[headers['姓名']];sid=values[headers['学号']]
                if not name and not sid:continue
                require(len(rows)<10000,'最多导入10000行',422)
                college=values[headers['学院']] if '学院' in headers else ''
                bad=any(cells[headers[k]].data_type=='f' for k in ('姓名','学号'))
                rows.append({'key':f'{sheet.title}:{number}','sheet':sheet.title,'row':number,'name':name,'student_id':sid,'college':college,'formula':bad});count+=1
            sheets.append({'name':sheet.title,'count':count,'group_name':SHEET_GROUPS.get(sheet.title),'business_name':'明远学术汇' if sheet.title=='明远学术' else '干部及干事' if sheet.title in SHEET_GROUPS else None})
        book.close();return raw,sheets,rows
    except DomainError:raise
    except Exception as e:raise DomainError('请上传有效且含姓名、学号表头的 .xlsx 文件',422) from e


class MemberImport:
    def import_members(self,actor,file_base64,mappings=None,decisions=None,commit=False,initial_password=None):
        from .service import require
        with self.db.transaction() as c:self.actor(c,actor,('SUPER',))
        raw,sheets,rows=parse_roster(file_base64);decisions=decisions or {};mappings=mappings or {}
        with self.db.transaction(True if commit else False) as c:
            a=self.actor(c,actor,('SUPER',));people={p['student_id']:p for p in all_rows(c,'SELECT student_id,name,college,status FROM student')}
            catalogue=all_rows(c,'SELECT * FROM business');groups=all_rows(c,'SELECT * FROM business_group')
            resolved={}
            for sheet in sheets:
                if sheet['name'] in mappings:target=mappings[sheet['name']]
                else:
                    b=next((b for b in catalogue if b['name']==sheet['business_name']),None)
                    g=next((g for g in groups if b and g['business_id']==b['business_id'] and g['name']==sheet['group_name']),None)
                    target={'business_id':b['business_id'],'group_id':g['group_id'] if g else None} if b else None
                if target:
                    b=self._active_business(c,int(target['business_id']));require(b['kind']=='EVENT','该入口只导入非固定业务／小组名单，固定业务请配置时段归属',422)
                    if target.get('group_id'):self._group(c,b['business_id'],int(target['group_id']))
                    target={'business_id':b['business_id'],'group_id':int(target['group_id']) if target.get('group_id') else None}
                resolved[sheet['name']]=target
            file_hash=hashlib.sha256(raw).hexdigest()
            request_hash=hashlib.sha256(json.dumps({'file':file_hash,'mappings':resolved,'decisions':decisions},sort_keys=True,ensure_ascii=False).encode()).hexdigest()
            old=one(c,'SELECT batch_id FROM member_import WHERE request_hash=%s',(request_hash,))
            if old:return {'replayed':True,'batch_id':old['batch_id'],'rows':[],'sheets':sheets,'errors':[],'new_students':0}
            errors=[];new={};seen={};eligible=[];skipped=0
            for r in rows:
                choice=decisions.get(r['key'],{});r['target']=resolved[r['sheet']];r['candidates']=[{'student_id':p['student_id'],'name':p['name'],'status':p['status']} for p in people.values() if p['name']==r['name']]
                if choice.get('skip') or r['target'] is None:r['skipped']=True;skipped+=1;continue
                sid=str(choice.get('student_id') or r['student_id']).strip();r['resolved_student_id']=sid
                problem=''
                if r['formula']:problem='姓名和学号不能使用公式'
                elif not sid or len(sid)>32 or not sid.isdigit():problem='缺少有效文本学号'
                elif sid in people and people[sid]['name']!=r['name']:problem='学号与系统姓名冲突，不覆盖原资料'
                elif sid in people and people[sid]['status']!='ACTIVE':problem='账号不在队，不自动启用'
                elif sid not in people and len(sid)!=13 and not choice.get('confirmed'):problem='学号长度与常用13位格式不同，请核对'
                elif sid not in people and r['candidates'] and not choice.get('confirmed'):problem='系统存在同名的其他学号，请核对身份'
                elif sid in seen and seen[sid]!=r['name']:problem='同一学号对应多个姓名'
                elif not r['name'] or len(r['name'])>100 or len(r['college'])>100:problem='姓名或学院长度无效'
                if problem:r['error']=problem;errors.append({'key':r['key'],'message':problem});continue
                seen[sid]=r['name'];eligible.append(r)
                if sid not in people:new[sid]=r
            result={'sheets':[{**s,'target':resolved[s['name']]} for s in sheets],'rows':rows,'errors':errors,'eligible':len(eligible),'skipped':skipped,'new_students':len(new)}
            if not commit:return result
            require(not errors,'仍有身份异常，请核对或明确跳过后再导入',422)
            require(eligible,'没有可导入的记录',422)
            require(not new or initial_password and 8<=len(initial_password)<=128,'请填写新账号初始密码',422)
            c.execute('INSERT INTO member_import(file_hash,request_hash,admin_id,created_at) VALUES(%s,%s,%s,%s)',(file_hash,request_hash,a['_id'],self.clock()));batch=c.lastrowid
            for sid,r in new.items():
                c.execute('INSERT INTO student(student_id,name,college,password_hash) VALUES(%s,%s,%s,%s)',(sid,r['name'],r['college'],hash_password(initial_password)))
            for r in eligible:
                sid=r['resolved_student_id'];target=r['target'];bid,gid=target['business_id'],target['group_id'];day=self.clock().date()
                self._join_business(c,bid,sid,day)
                if gid:self._join_group(c,gid,sid,day)
                c.execute('INSERT INTO member_import_row VALUES(%s,%s,%s,%s,%s,%s)',(batch,r['sheet'],r['row'],sid,bid,gid))
            self.audit(c,a['_id'],'IMPORT_MEMBERS','member_import',batch,after={'imported':len(eligible),'skipped':skipped,'new_students':len(new)},reason='名单导入；只建立账号和归属，不计时')
            return {**result,'batch_id':batch,'imported':len(eligible),'new_accounts':[{'student_id':sid,'name':r['name']} for sid,r in new.items()]}

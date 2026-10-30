"""Read-only XLSX preview followed by a transactional, replay-safe import."""
import base64
import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile
from openpyxl import load_workbook
from .db import one, all_rows
from .security import hash_password


HEADERS=['学号','姓名','学院','服务日期','工作事项','认定时长（小时）','备注']


def parse_specialty(encoded):
    from .service import require,DomainError
    try:
        raw=base64.b64decode(encoded,validate=True)
        require(len(raw)<=8*1024*1024,'文件不能超过8MB',422)
        with ZipFile(BytesIO(raw)) as z:
            require(len(z.infolist())<=2000 and sum(x.file_size for x in z.infolist())<=40*1024*1024,'工作簿解压内容过大',422)
        book=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    except DomainError:raise
    except Exception as e:raise DomainError('请上传有效的 .xlsx 文件',422) from e
    rows=[];errors=[];seen=set();people={}
    try:
        sheet=book.worksheets[0]
        require((sheet.max_row or 0)<=10001 and (sheet.max_column or 0)<=30,'最多导入10000行、30列',422)
        iterator=sheet.iter_rows();header={str(cell.value).strip():i for i,cell in enumerate(next(iterator,())) if cell.value is not None}
        require(set(HEADERS)-{'学院','备注'}<=set(header),'第一行须包含模板的必填表头',422)
        for n,row in enumerate(iterator,2):
            require(n<=10001,'最多导入10000行',422)
            if not any(c.value is not None for c in row):continue
            try:
                values={key:row[idx].value for key,idx in header.items() if key in HEADERS}
                require(not any(row[idx].data_type=='f' for key,idx in header.items() if key in HEADERS),'不接受公式，请改为明确的数值或文本',422)
                sid=values['学号'];name=str(values['姓名'] or '').strip()
                require(isinstance(sid,str) and re.fullmatch(r'[A-Za-z0-9_-]{1,32}',sid.strip()),'学号必须为文本，且只能包含字母、数字、下划线或短横线',422)
                sid=sid.strip();require(1<=len(name)<=100,'姓名不能为空或超过100字',422)
                require(sid not in people or people[sid]==name,'同一学号出现不同姓名',422);people[sid]=name
                day=values['服务日期']
                if isinstance(day,datetime):day=day.date()
                elif not isinstance(day,date):day=date.fromisoformat(str(day).strip())
                title=str(values['工作事项'] or '').strip();require(1<=len(title)<=100,'工作事项须为1至100字',422)
                duration=Decimal(str(values['认定时长（小时）']))*60
                require(duration.is_finite() and duration==duration.to_integral_value() and 0<=duration<=1440,'时长必须准确换算为0至1440的整数分钟',422)
                key=(sid,day.isoformat(),title,int(duration))
                require(key not in seen,'同表存在重复记录，请合并或删除',422);seen.add(key)
                college=str(values.get('学院') or '').strip();remark=str(values.get('备注') or '').strip()
                require(len(college)<=100 and len(remark)<=500,'学院或备注过长',422)
                rows.append({'row':n,'student_id':sid,'name':name,'college':college,'service_date':day,'title':title,'service_minutes':int(duration),'remark':remark})
            except (DomainError,ValueError,InvalidOperation,TypeError,KeyError) as e:
                errors.append({'row':n,'message':str(e) if isinstance(e,DomainError) else '日期或时长格式无效'})
    finally:book.close()
    return raw,rows,errors


class SpecialtyService:
    def import_specialty(self,actor,business_name,file_base64,commit=False,initial_password=None,duplicate_rows=None):
        from .service import require
        with self.db.transaction() as c:self.actor(c,actor,('SUPER',))
        raw,rows,errors=parse_specialty(file_base64)
        name=business_name.strip();require(1<=len(name)<=100,'请填写专项业务名称',422)
        digest=hashlib.sha256(raw).hexdigest()
        choices={int(k):v for k,v in (duplicate_rows or {}).items()}
        require(all(v in ('skip','separate') for v in choices.values()),'重复行处理方式无效',422)
        with self.db.transaction(commit) as c:
            a=self.actor(c,actor,('SUPER',));biz=one(c,'SELECT * FROM business WHERE name=%s',(name,))
            require(not biz or biz['kind']=='EVENT','该名称已用于固定业务，请选择专项业务名称',422)
            prior=one(c,'SELECT batch_id FROM specialty_import WHERE business_id=%s AND file_hash=%s',(biz['business_id'],digest)) if biz else None
            if prior:return {'batch_id':prior['batch_id'],'replayed':True,'rows':[],'errors':[],'duplicates':[],'new_students':0}
            people={r['student_id']:r for r in all_rows(c,'SELECT student_id,name,status FROM student')};duplicates=[]
            for r in rows:
                p=people.get(r['student_id'])
                if p and (p['name']!=r['name'] or p['status']!='ACTIVE'):errors.append({'row':r['row'],'message':'学号已存在但姓名不同或已停用/退队，不会覆盖'})
                if r['service_date']>self.clock().date():errors.append({'row':r['row'],'message':'服务日期不能为未来日期'})
                r['fingerprint']=hashlib.sha256(json.dumps([name,r['student_id'],str(r['service_date']),r['title']],ensure_ascii=False).encode()).hexdigest()
                match=one(c,'SELECT 1 FROM specialty_import_row WHERE fingerprint=%s LIMIT 1',(r['fingerprint'],))
                if biz and not match:
                    match=one(c,'''SELECT 1 FROM schedule s JOIN service_slot t ON t.slot_id=s.slot_id
                        WHERE t.business_id=%s AND s.student_id=%s AND s.schedule_date=%s AND t.team_name=%s
                        AND s.status IN ('ACTIVE','COMPLETED') LIMIT 1''',(biz['business_id'],r['student_id'],r['service_date'],r['title']))
                if match:duplicates.append(r['row'])
            new={r['student_id'] for r in rows if r['student_id'] not in people}
            preview={'business_name':name,'rows':rows,'errors':errors,'duplicates':duplicates,'new_students':len(new),'file_hash':digest}
            if not commit:return preview
            require(not errors,'请先修正预览中的错误',422);require(rows,'工作簿没有数据',422)
            require(all(n in choices for n in duplicates),'请逐行确认疑似重复记录是跳过还是另一项独立服务',422)
            selected=[r for r in rows if r['row'] not in duplicates or choices[r['row']]=='separate']
            new={r['student_id'] for r in selected if r['student_id'] not in people}
            require(not new or initial_password and 8<=len(initial_password)<=128,'请设置8至128字符的新队员初始密码',422)
            biz=self._event_business(c,name)
            c.execute('INSERT INTO specialty_import(business_id,file_hash,admin_id,create_time) VALUES(%s,%s,%s,%s)',(biz['business_id'],digest,a['_id'],self.clock()));batch=c.lastrowid
            encoded=hash_password(initial_password) if new else None
            for r in selected:
                sid=r['student_id']
                if sid not in people:
                    c.execute('INSERT INTO student(student_id,name,college,password_hash) VALUES(%s,%s,%s,%s)',(sid,r['name'],r['college'],encoded));people[sid]=r
                schedule=self._event_record(c,sid,biz['business_id'],r['title'],r['service_date'],r['service_minutes'],a['_id'],'SPECIALTY_IMPORT',f'import:{batch}:{r["row"]}')
                c.execute('UPDATE attendance SET remark=%s WHERE schedule_id=%s',(r['remark'],schedule))
                c.execute('INSERT INTO specialty_import_row(batch_id,row_no,fingerprint,schedule_id) VALUES(%s,%s,%s,%s)',(batch,r['row'],r['fingerprint'],schedule))
            result={'batch_id':batch,'imported':len(selected),'skipped':len(rows)-len(selected),'new_students':len(new)}
            self.audit(c,a['_id'],'IMPORT_SPECIALTY','specialty_import',batch,after={**result,'business_name':name,'duplicate_choices':choices},reason='专项名单及时长导入')
            return result

"""Integration coverage against volunteer_core_test, never the demo database."""
import base64
import json
import sys
import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

CORE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(CORE))
from app.bootstrap import initialize_database
from app.fixed_slots import migrate
from app.db import Database,one,all_rows
from app.service import Service,DomainError
from app.specialty import parse_specialty


class BusinessTests(unittest.TestCase):
    def setUp(self):
        cfg=json.loads((CORE.parent/'.work/mysql-core/admin.json').read_text(encoding='utf-8'))
        self.demo=initialize_database(cfg,database='volunteer_core_test',reset=True)
        self.settings={**cfg,'database':'volunteer_core_test'};migrate(self.settings)
        self.db=Database(self.settings);self.now=datetime(2026,9,6,7)
        self.s=Service(self.db,clock=lambda:self.now)
        self.sup,self.normal,self.other,self.student=[self.demo[k] for k in ('super','normal','other_admin','student')]
        self.sid=self.student['id']
        self.b={b['code'] or b['name']:b for b in self.s.businesses(self.sup)}
        self.b['干部']=self.b['干部及干事']
        self.slots=self.s.list_slots(self.sup)
        self.consult=next(s['slot_id'] for s in self.slots if s['business_code']=='CONSULT' and s['weekday']==1 and s['period_no']==1)
        self.art=next(s['slot_id'] for s in self.slots if s['business_code']=='ART' and s['weekday']==1 and s['period_no']==3)
        self.pub=next(s['slot_id'] for s in self.slots if s['business_code']=='PUBLICITY' and s['weekday']==1 and s['period_no']==1)
        self.garden=[s['slot_id'] for s in self.slots if s['business_code']=='GARDEN']
        self.s.configure_slot(self.sup,self.consult,self.normal['id'],20)
        self.s.configure_slot(self.sup,self.art,self.other['id'],20)
        for code in ('GARDEN','PUBLICITY'):self.s.configure_business(self.sup,self.b[code]['business_id'],self.normal['id'])

    def rows(self,sql,args=()):
        with self.db.transaction() as c:return list(all_rows(c,sql,args))

    def term(self,slots):
        self.s.set_memberships(self.sup,self.sid,slots)
        return self.s.save_semester(self.sup,'多业务测试','2026-09-07',4)['semester_id']

    def test_catalog_migration_preserves_ledger_and_ids(self):
        self.term([self.consult]);self.now=datetime(2026,9,14);self.s.run_weekly_settlement()
        before=self.rows('SELECT * FROM service_record');ids=[s['slot_id'] for s in self.s.list_slots(self.sup)]
        migrate(self.settings)
        self.assertEqual(before,self.rows('SELECT * FROM service_record'))
        self.assertEqual(ids,[s['slot_id'] for s in self.s.list_slots(self.sup)])
        for code,count in [('CONSULT',35),('ART',28),('PUBLICITY',21),('GARDEN',2)]:
            self.assertEqual(sum(s['business_code']==code and s['status']=='ACTIVE' for s in self.slots),count)
        self.assertEqual(next(r for r in self.s.rules(self.sup) if r['status']=='ACTIVE' and r['violation_type']=='LATE')['deduct_score'],1)

    def test_multi_weekly_totals_and_scope(self):
        self.term([self.consult,self.pub,self.garden[0]])
        self.now=datetime(2026,9,14);self.assertEqual(self.s.run_weekly_settlement()['settled'],3)
        self.assertEqual(self.s.hours(self.student)['total_minutes'],390)
        self.assertEqual(self.s.run_weekly_settlement()['settled'],0)
        self.assertEqual(len(self.s.list_schedules(self.other)),0)
        self.assertEqual(self.s.hours(self.student,business_id=self.b['GARDEN']['business_id'])['total_minutes'],90)

    def test_garden_midweek_confirm_and_ab_change_not_double(self):
        self.term([]);self.now=datetime(2026,9,9,12)
        self.s.set_memberships(self.sup,self.sid,[self.garden[0]])
        row=self.s.list_schedules(self.student)[0];self.assertEqual(row['weekly_confirmed'],0)
        with self.assertRaises(DomainError):self.s.confirm_garden(self.other,row['schedule_id'],'不属于我')
        self.s.confirm_garden(self.normal,row['schedule_id'],'本周已完成')
        self.s.set_memberships(self.sup,self.sid,[self.garden[1]])
        self.now=datetime(2026,9,14);self.s.run_weekly_settlement()
        self.assertEqual(self.s.hours(self.student)['total_minutes'],90)
        self.now=datetime(2026,9,21);self.s.run_weekly_settlement()
        self.assertEqual(self.s.hours(self.student)['total_minutes'],180)

    def test_unconfirmed_join_not_auto_awarded_and_no_garden_late(self):
        self.term([]);self.now=datetime(2026,9,9)
        self.s.set_memberships(self.sup,self.sid,[self.garden[0]])
        row=self.s.list_schedules(self.student)[0];self.now=datetime(2026,9,14)
        self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.student)['total_minutes'],0)
        with self.assertRaises(DomainError):self.s.record_attendance(self.normal,row['schedule_id'],'LATE','')
        self.s.confirm_garden(self.normal,row['schedule_id'],'补充确认已完成')
        self.assertEqual(self.s.hours(self.student)['total_minutes'],90)

    def test_garden_current_week_absence_can_be_corrected_without_fake_confirmation(self):
        self.term([]);self.now=datetime(2026,9,9)
        self.s.set_memberships(self.sup,self.sid,[self.garden[0]])
        row=self.s.list_schedules(self.student)[0]
        self.s.record_attendance(self.normal,row['schedule_id'],'ABSENT','本周未完成')
        self.assertEqual(self.s.credit(self.student)['credit_score'],2)
        self.s.correct_attendance(self.normal,row['schedule_id'],'NORMAL','核实本周已完成',1)
        self.assertEqual(self.s.credit(self.student)['credit_score'],5)
        self.assertEqual(self.s.hours(self.student)['total_minutes'],90)

    def test_stop_restore_hours_credit_and_history_permission(self):
        term=self.term([self.pub,self.garden[0]]);self.now=datetime(2026,9,14)
        self.s.run_weekly_settlement();row=next(r for r in self.s.list_schedules(self.student) if r['slot_id']==self.pub and r['week_no']==1)
        self.s.correct_attendance(self.normal,row['schedule_id'],'ABSENT','补录缺勤',1)
        self.assertEqual(self.s.credit(self.student)['credit_score'],2)
        with self.assertRaises(DomainError):self.s.set_business_week(self.normal,self.b['PUBLICITY']['business_id'],term,1,False,'历史周停工')
        self.s.set_business_week(self.sup,self.b['PUBLICITY']['business_id'],term,1,False,'核实原定停工')
        self.assertEqual(self.s.credit(self.student)['credit_score'],5)
        self.s.set_business_week(self.sup,self.b['PUBLICITY']['business_id'],term,1,True,'恢复原安排')
        self.assertEqual(self.s.credit(self.student)['credit_score'],2)
        self.assertEqual(self.s.hours(self.student)['total_minutes'],90)
        self.assertTrue(self.s.set_business_week(self.sup,self.b['PUBLICITY']['business_id'],term,1,True,'重试')['replayed'])

    def test_stopped_week_no_auto_hours_and_restore_manual_minutes(self):
        term=self.term([self.garden[0]]);self.now=datetime(2026,9,8)
        self.s.set_business_week(self.normal,self.b['GARDEN']['business_id'],term,1,False,'本周暂停')
        self.now=datetime(2026,9,14);self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.student)['total_minutes'],0)
        self.s.set_business_week(self.sup,self.b['GARDEN']['business_id'],term,1,True,'本周确实完成')
        record=self.s.hours(self.student)['records'][0];self.s.revise_hours(self.sup,record['record_id'],80,'核实80分钟')
        self.s.set_business_week(self.sup,self.b['GARDEN']['business_id'],term,1,False,'再次核实')
        self.s.set_business_week(self.sup,self.b['GARDEN']['business_id'],term,1,True,'保留人工认定')
        self.assertEqual(self.s.hours(self.student)['total_minutes'],80)

    def test_exclusive_effective_membership_keeps_history(self):
        self.term([self.consult]);
        with self.assertRaises(DomainError):self.s.set_memberships(self.sup,self.sid,[self.consult,self.art])
        self.s.set_memberships(self.sup,self.sid,[self.art],'2026-09-14','转入美育')
        self.now=datetime(2026,9,22);self.s.run_weekly_settlement()
        hours=self.s.hours(self.student);self.assertEqual(hours['total_minutes'],300)
        self.assertEqual({x['business_name'] for x in hours['business_totals']},{'咨询服务','美育'})

    def test_claim_range_per_day_return_and_repeat_approval(self):
        self.term([]);self.now=datetime(2026,9,14)
        cid=self.s.create_collection(self.sup,self.b['干部']['business_id'],'学期筹备','2026-09-07','2026-09-20',[self.sid],[{'title':'筹备','min_minutes':120,'max_minutes':240}])['collection_id']
        item=self.s.collections(self.student)[0]['items'][0]['item_id']
        with self.assertRaises(DomainError):self.s.submit_hours(self.demo['other_student'],item,'2026-09-07',120)
        with self.assertRaises(DomainError):self.s.submit_hours(self.student,item,'2026-09-07',241)
        oneid=self.s.submit_hours(self.student,item,'2026-09-07',240)['submission_id']
        twoid=self.s.submit_hours(self.student,item,'2026-09-08',240)['submission_id']
        self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.student)['total_minutes'],0)
        self.s.decide_hours(self.sup,oneid,False,1,'请补充说明')
        self.s.submit_hours(self.student,item,'2026-09-07',180,'完成会场筹备',2)
        self.s.decide_hours(self.sup,oneid,True,3)
        self.s.decide_hours(self.sup,twoid,True,1)
        self.assertTrue(self.s.decide_hours(self.sup,twoid,True,1)['replayed'])
        self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.student)['total_minutes'],420)

    def test_event_ledger_audit_rollback_and_export(self):
        term=self.term([]);self.now=datetime(2026,9,14)
        with patch.object(self.s,'audit',side_effect=RuntimeError('audit failed')):
            with self.assertRaises(RuntimeError):self.s.create_event(self.sup,'测试专项',self.sid,'现场服务','2026-09-08',123,'实际完成','x'*16)
        self.assertEqual(self.s.hours(self.student)['total_minutes'],0)
        sid=self.s.create_event(self.sup,'测试专项',self.sid,'现场服务','2026-09-08',123,'实际完成','x'*16)['schedule_id']
        self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.student)['total_minutes'],123)
        with self.assertRaises(DomainError):self.s.correct_attendance(self.sup,sid,'ABSENT','不应产生专项扣分',1)
        from app.hours_export import export_hours
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(export_hours(self.s,self.sup,term)))
        matches=[r for r in book.worksheets[0].iter_rows(min_row=7,values_only=True) if r[1]=='测试专项']
        self.assertEqual(matches[0][5],2.05)

    def test_leave_image_scope_and_rejection_reason(self):
        self.term([self.consult])
        from PIL import Image
        output=BytesIO();Image.new('RGB',(4,4),'white').save(output,format='PNG');img=base64.b64encode(output.getvalue()).decode()
        row=self.s.list_schedules(self.student)[0]
        with self.assertRaises(DomainError):self.s.apply_leave(self.student,row['schedule_id'],'请假')
        lid=self.s.apply_leave(self.student,row['schedule_id'],'请假',[img])['leave_id']
        self.assertEqual(len(self.s.leave_evidence(self.normal,lid)),1)
        with self.assertRaises(DomainError):self.s.leave_evidence(self.other,lid)
        with self.assertRaises(DomainError):self.s.approve_leave(self.normal,lid,False)
        self.s.approve_leave(self.normal,lid,False,'证明不完整')
        self.assertEqual(self.s.list_applications(self.student,'leave')[0]['rejection_reason'],'证明不完整')

    def test_specialty_preview_import_repeat_and_explicit_duplicates(self):
        self.now=datetime(2026,9,21)
        data=(CORE/'templates/专项时长导入示例.xlsx').read_bytes();encoded=base64.b64encode(data).decode()
        preview=self.s.import_specialty(self.sup,'新建专项',encoded)
        self.assertEqual(preview['new_students'],2);self.assertEqual(preview['errors'],[])
        self.assertFalse(self.rows("SELECT 1 FROM business WHERE name='新建专项'"))
        with self.assertRaises(DomainError):self.s.import_specialty(self.normal,'新建专项',encoded)
        result=self.s.import_specialty(self.sup,'新建专项',encoded,True,'TestPassword123!')
        self.assertEqual(result['imported'],3)
        self.assertTrue(self.s.import_specialty(self.sup,'新建专项',encoded,True)['replayed'])
        self.assertEqual(self.s.hours(self.sup,'0012026001')['total_minutes'],210)
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(data));book.active['G2']='另一份批次需核对重复';buf=BytesIO();book.save(buf)
        changed=base64.b64encode(buf.getvalue()).decode()
        self.assertEqual(self.s.import_specialty(self.sup,'新建专项',changed)['duplicates'],[2,3,4])
        with self.assertRaises(DomainError):self.s.import_specialty(self.sup,'新建专项',changed,True)
        self.s.import_specialty(self.sup,'新建专项',changed,True,duplicate_rows={2:'separate',3:'skip',4:'skip'})
        self.assertEqual(self.s.hours(self.sup,'0012026001')['total_minutes'],330)
        self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.sup,'0012026001')['total_minutes'],330)

    def test_specialty_invalid_data_and_atomic_audit(self):
        self.now=datetime(2026,9,21)
        data=(CORE/'templates/专项时长导入示例.xlsx').read_bytes();encoded=base64.b64encode(data).decode()
        with patch.object(self.s,'audit',side_effect=RuntimeError('audit failed')):
            with self.assertRaises(RuntimeError):self.s.import_specialty(self.sup,'回滚专项',encoded,True,'TestPassword123!')
        self.assertFalse(self.rows("SELECT 1 FROM student WHERE student_id='0012026001'"))
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(data));book.active['F2']=0.001;book.active['A3']=123;book.active['F4']='=1+1'
        buf=BytesIO();book.save(buf)
        preview=self.s.import_specialty(self.sup,'错误专项',base64.b64encode(buf.getvalue()).decode())
        self.assertEqual(len(preview['errors']),3)

    def test_future_event_requires_explicit_recognition_and_membership_only_has_no_hours(self):
        self.term([])
        sid=self.s.create_event(self.sup,'讲解队',self.sid,'临时讲解','2026-09-08',120,'预先安排','z'*16)['schedule_id']
        with self.assertRaises(DomainError):self.s.recognize_event(self.sup,sid,90,'完成确认')
        self.now=datetime(2026,9,14);self.s.run_weekly_settlement()
        self.assertEqual(self.s.hours(self.student)['total_minutes'],0)
        with self.assertRaises(DomainError):self.s.recognize_event(self.normal,sid,90,'越权')
        self.s.recognize_event(self.sup,sid,90,'实际认定90分钟')
        self.assertTrue(self.s.recognize_event(self.sup,sid,90,'重试')['replayed'])
        self.s.set_event_member(self.sup,self.b['干部']['business_id'],self.sid,True,'2026-09-14','加入干部业务')
        self.assertIn(self.b['干部']['business_id'],self.s.students(self.sup,self.sid)[0]['business_ids'])
        self.assertEqual(self.s.hours(self.student)['total_minutes'],90)

    def test_fixed_business_roster_import_checks_business_period_and_exclusivity(self):
        from test_roster import fixture
        book=fixture([['张三',self.sid,'','一分队三时段']])
        self.s.import_roster(self.sup,book,commit=True,initial_password='TestPassword123!',business_id=self.b['ART']['business_id'])
        self.assertEqual(self.s.my_memberships(self.student)[0]['slot_id'],self.art)
        self.assertTrue(self.s.import_roster(self.sup,book)['errors'])
        bad=fixture([['李四',self.demo['other_student']['id'],'','一分队五时段']])
        self.assertTrue(self.s.import_roster(self.sup,bad,business_id=self.b['ART']['business_id'])['errors'])

    def test_future_change_catches_up_after_term_ended(self):
        self.term([self.consult])
        self.s.set_memberships(self.sup,self.sid,[self.art],'2026-09-14','第二周开始转美育')
        self.now=datetime(2026,10,12)
        self.s.run_weekly_settlement()
        self.assertEqual(self.s.hours(self.student,business_id=self.b['CONSULT']['business_id'])['total_minutes'],120)
        self.assertEqual(self.s.hours(self.student,business_id=self.b['ART']['business_id'])['total_minutes'],540)

    def test_new_api_routes_authorization_and_template(self):
        import os
        os.environ['VOLUNTEER_DISABLE_SCHEDULER']='1'
        from fastapi.testclient import TestClient
        from app.api import app,get_service,identity
        old=dict(app.dependency_overrides)
        app.dependency_overrides[get_service]=lambda:self.s
        app.dependency_overrides[identity]=lambda:self.sup
        try:
            with TestClient(app) as client:
                self.assertEqual(client.get('/businesses').status_code,200)
                self.assertEqual(client.get('/specialty-imports/template').status_code,200)
                response=client.post('/hour-collections',json={'business_id':self.b['干部']['business_id'],'title':'测试征集','start_date':'2026-09-07','end_date':'2026-09-20','student_ids':[self.sid],'items':[{'title':'准备','min_minutes':60,'max_minutes':180}]})
                self.assertEqual(response.status_code,200,response.text)
                app.dependency_overrides[identity]=lambda:self.normal
                self.assertEqual(client.get('/hour-collections').status_code,200)
                self.assertEqual(client.get('/hour-collections').json(),[])
                self.assertEqual(client.get('/specialty-imports/template').status_code,403)
                self.assertEqual(client.put(f'/businesses/{self.b["GARDEN"]["business_id"]}/manager',json={'manager_admin_id':self.normal['id']}).status_code,403)
        finally:app.dependency_overrides.clear();app.dependency_overrides.update(old)


if __name__=='__main__':unittest.main()

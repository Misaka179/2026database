"""Weekly report calculations and real DB permission/filtering regression tests."""
import sys
import unittest
from pathlib import Path
from datetime import date, datetime
from io import BytesIO
from openpyxl import load_workbook

CORE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(CORE))
from app.hours_export import build_workbook, export_hours
from app.service import DomainError


class WorkbookTests(unittest.TestCase):
    def test_week_dates_multi_slot_quit_and_future(self):
        term={'name':'测试学期','start_date':date(2026,9,2),'end_date':date(2026,9,20),'end_week':3}
        students=[{'student_id':'001','name':'=1+1','college':'学院','status':'ACTIVE'},
                  {'student_id':'002','name':'退队者','college':'','status':'QUIT'}]
        slots=[{'slot_id':4,'weekday':1,'period_no':2},{'slot_id':7,'weekday':2,'period_no':3}]
        memberships=[{'student_id':'001','slot_id':4},{'student_id':'001','slot_id':7}]
        records=[{'student_id':'001','slot_id':4,'schedule_date':date(2026,9,3),'service_minutes':120},
                 {'student_id':'001','slot_id':7,'schedule_date':date(2026,9,4),'service_minutes':180},
                 {'student_id':'002','slot_id':4,'schedule_date':date(2026,9,4),'service_minutes':75},
                 {'student_id':'001','slot_id':4,'schedule_date':date(2026,8,31),'service_minutes':120}]
        book=load_workbook(BytesIO(build_workbook(term,students,slots,memberships,records,datetime(2026,9,8))))
        main,summary,quit_sheet=book.worksheets
        self.assertEqual(book.sheetnames,['志愿时长统计','分队时长统计','退队统计'])
        self.assertEqual(main['F4'].value,'第 1 周\n2026-08-31～2026-09-06')
        self.assertEqual(main['E7'].value,'001')
        self.assertEqual(main['E7'].data_type,'s')
        self.assertEqual(main['A7'].data_type,'s')
        self.assertEqual([main.cell(r,6).value for r in (7,8)],[2,3])
        self.assertEqual([main.cell(r,9).value for r in (7,8)],[2,3])
        self.assertEqual(main['G7'].value,0)
        self.assertIsNone(main['H7'].value)
        self.assertEqual(quit_sheet['F7'].value,1.25)
        self.assertEqual(summary['B14'].value,6.25)
        self.assertEqual(summary['E14'].value,6.25)
        self.assertEqual(main.freeze_panes,'F7')


class ExportDatabaseTests(unittest.TestCase):
    def setUp(self):
        # Reuse the established isolated DB fixture; never write production data.
        import test_semester
        self.fixture=test_semester.SemesterTests()
        self.fixture.setUp()
        self.term=self.fixture.term()

    def test_only_super_can_export_and_invalid_term_is_rejected(self):
        f=self.fixture
        for actor in (f.normal,f.student):
            with self.assertRaises(DomainError) as ex:export_hours(f.s,actor,self.term)
            self.assertEqual(ex.exception.status,403)
        with self.assertRaises(DomainError) as ex:export_hours(f.s,f.super,999999)
        self.assertEqual(ex.exception.status,404)

    def test_settlement_absence_correction_and_manual_minutes_are_exported(self):
        f=self.fixture
        f.now=datetime(2026,9,14,9)
        f.s.run_weekly_settlement()
        initial=load_workbook(BytesIO(export_hours(f.s,f.super,self.term)))
        total=lambda b:b['分队时长统计'].cell(14,5).value
        self.assertEqual(total(initial),7)
        sid=next(r['schedule_id'] for r in f.s.list_schedules(f.student)
                 if str(r['schedule_date'])=='2026-09-07' and r['slot_id']==f.mon)
        with f.db.transaction() as c:
            from app.db import one
            attendance=one(c,'SELECT revision FROM attendance WHERE schedule_id=%s',(sid,))
        f.s.correct_attendance(f.super,sid,'ABSENT','纠正缺勤',attendance['revision'])
        after=load_workbook(BytesIO(export_hours(f.s,f.super,self.term)))
        self.assertEqual(total(after),5)
        with f.db.transaction() as c:
            record=one(c,'''SELECT r.record_id FROM service_record r
                JOIN attendance a ON a.attendance_id=r.attendance_id
                JOIN schedule s ON s.schedule_id=a.schedule_id
                WHERE s.student_id=%s AND s.slot_id=%s AND s.schedule_date='2026-09-08' ''',
                (f.student['id'],f.tue))
        f.s.revise_hours(f.super,record['record_id'],135,'修正实际认定分钟')
        corrected=load_workbook(BytesIO(export_hours(f.s,f.super,self.term)))
        self.assertEqual(total(corrected),4.25)
        from fastapi.testclient import TestClient
        from app.api import app,identity,get_service
        old=dict(app.dependency_overrides)
        app.dependency_overrides[identity]=lambda:f.super
        app.dependency_overrides[get_service]=lambda:f.s
        try:
            with TestClient(app) as client:
                response=client.get('/hours/export',params={'semester_id':self.term})
                self.assertEqual(response.status_code,200)
                self.assertIn('attachment',response.headers['content-disposition'])
                self.assertEqual(total(load_workbook(BytesIO(response.content))),4.25)
        finally:
            app.dependency_overrides.clear();app.dependency_overrides.update(old)


if __name__=='__main__':unittest.main()

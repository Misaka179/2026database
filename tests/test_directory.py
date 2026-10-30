"""Isolated lifecycle, group authorization and roster compatibility tests."""
import base64,json,unittest
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook,load_workbook
import test_business
from app.service import DomainError
from app.fixed_slots import migrate
from app.hours_export import export_hours


class DirectoryTests(unittest.TestCase):
    setUp=test_business.BusinessTests.setUp
    rows=test_business.BusinessTests.rows
    term=test_business.BusinessTests.term

    def slot(self,p=1):
        return {'weekday':1,'period_no':p,'start_time':f'{p+6:02d}:00','end_time':f'{p+7:02d}:00','service_minutes':60,'need_num':20,'manager_admin_id':int(self.normal['id'])}

    def group(self,n=0):return self.s.business_groups(self.sup)[n]['group_id']

    def test_custom_six_periods_no_backfill_and_export(self):
        term=self.s.save_semester(self.sup,'秋季','2026-09-07',4)['semester_id']
        self.now=datetime(2026,9,15,8)
        bid=self.s.create_business(self.sup,'自定义业务','FIXED',[self.slot(p) for p in range(1,7)])['business_id']
        slots=self.rows('SELECT * FROM service_slot WHERE business_id=%s ORDER BY period_no',(bid,));self.assertEqual(len(slots),6)
        self.s.set_memberships(self.sup,self.sid,[slots[-1]['slot_id']])
        self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.student)['total_minutes'],0)
        self.now=datetime(2026,9,28);self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.student)['total_minutes'],60)
        self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.student)['total_minutes'],60)
        book=load_workbook(BytesIO(export_hours(self.s,self.sup,term,bid)),data_only=True)
        self.assertEqual(len(book.sheetnames),3);self.assertTrue(any('6时段' in str(c.value) for row in book.worksheets[0] for c in row))
        with self.assertRaises(DomainError):self.s.add_business_slot(self.sup,bid,**self.slot(6))
        with self.assertRaises(DomainError):self.s.create_business(self.normal,'越权','EVENT')

    def test_archive_restore_gap_history_and_startup_idempotence(self):
        bid=self.s.create_business(self.sup,'停用测试','FIXED',[self.slot()])['business_id'];sid=self.rows('SELECT slot_id FROM service_slot WHERE business_id=%s',(bid,))[0]['slot_id']
        self.term([sid]);self.now=datetime(2026,9,14,6);self.s.run_weekly_settlement()
        before=self.s.hours(self.student)['total_minutes'];self.assertEqual(before,60)
        impact=self.s.archive_impact(self.sup,bid);self.assertGreater(impact['future_count'],0)
        self.s.change_business_status(self.sup,bid,False,1,'结束本阶段');migrate(self.settings)
        self.assertEqual(self.rows('SELECT status FROM business WHERE business_id=%s',(bid,))[0]['status'],'ARCHIVED')
        self.now=datetime(2026,9,22);self.s.run_weekly_settlement();self.assertEqual(self.s.hours(self.student)['total_minutes'],before)
        with self.assertRaises(DomainError):self.s.create_schedule(self.sup,self.sid,sid,'2026-09-28')
        self.s.change_business_status(self.sup,bid,True,2,'重新开展');self.now=datetime(2026,10,5);self.s.run_weekly_settlement()
        self.assertEqual(self.s.hours(self.student)['total_minutes'],120)
        self.assertEqual(len(self.rows("SELECT * FROM schedule WHERE slot_id=%s AND schedule_date IN ('2026-09-14','2026-09-21') AND status<>'CANCELLED'",(sid,))),0)

    def test_archive_blocked_by_collection_and_pending_claim(self):
        bid=self.b['干部']['business_id'];cid=self.s.create_collection(self.sup,bid,'开放征集','2026-09-01','2026-09-06',[self.sid],[{'title':'工作','min_minutes':30,'max_minutes':60}])['collection_id']
        self.assertTrue(self.s.archive_impact(self.sup,bid)['blockers'])
        with self.assertRaises(DomainError):self.s.change_business_status(self.sup,bid,False,1,'停用')
        item=self.rows('SELECT item_id FROM collection_item WHERE collection_id=%s',(cid,))[0]['item_id'];sub=self.s.submit_hours(self.student,item,'2026-09-06',30)['submission_id']
        self.s.close_collection(self.sup,cid,'停止收集')
        with self.assertRaises(DomainError):self.s.change_business_status(self.sup,bid,False,1,'停用')
        self.s.decide_hours(self.sup,sub,True,1);self.s.change_business_status(self.sup,bid,False,1,'停用')
        with self.assertRaises(DomainError):self.s.create_event(self.sup,'干部',self.sid,'任务','2026-09-06',30,'记录','a'*16)
        self.assertEqual(self.s.hours(self.student)['total_minutes'],30)

    def test_archive_pending_assignment_garden_week_and_event_restore(self):
        bid=self.b['GARDEN']['business_id'];self.term([self.garden[0]])
        self.now=datetime(2026,9,9,12);self.s.change_business_status(self.sup,bid,False,1,'暂时停止')
        self.now=datetime(2026,9,21);self.s.run_weekly_settlement()
        self.assertEqual(self.s.hours(self.student,business_id=bid)['total_minutes'],90)
        self.s.set_memberships(self.sup,self.sid,[self.consult],effective_date='2026-09-28',reason='待生效')
        self.assertTrue(any(r['kind']=='待生效归属调整' for r in self.s.archive_impact(self.sup,self.b['CONSULT']['business_id'])['blockers']))
        event=self.s.create_business(self.sup,'未来事项','EVENT')['business_id']
        first=self.s.create_event(self.sup,'未来事项',self.sid,'一','2026-10-01',60,'测试','future-one-key1')['schedule_id']
        second=self.s.create_event(self.sup,'未来事项',self.sid,'二','2026-10-02',60,'测试','future-two-key2')['schedule_id']
        with self.db.transaction(True) as c:c.execute("UPDATE schedule SET status='CANCELLED',cancel_reason='MANUAL' WHERE schedule_id=%s",(second,))
        self.s.change_business_status(self.sup,event,False,1,'停用');self.s.change_business_status(self.sup,event,True,2,'恢复')
        self.assertEqual(self.rows('SELECT status FROM schedule WHERE schedule_id=%s',(first,))[0]['status'],'ACTIVE')
        self.assertEqual(self.rows('SELECT cancel_reason FROM schedule WHERE schedule_id=%s',(second,))[0]['cancel_reason'],'MANUAL')

    def test_group_scope_files_membership_snapshot_and_totals(self):
        bid=self.b['干部']['business_id'];gid=self.group();other=self.group(1)
        self.s.set_group_member(self.sup,gid,self.sid,True,'2026-09-06','加入')
        self.s.set_group_member(self.sup,other,self.sid,True,'2026-09-06','兼任')
        self.s.set_collection_managers(self.sup,self.normal['id'],[],'小组负责人',[gid])
        self.s.set_collection_managers(self.sup,self.other['id'],[],'另一小组负责人',[other])
        f={'filename':'说明.txt','file_base64':base64.b64encode('说明'.encode()).decode()}
        cid=self.s.create_collection(self.normal,bid,'小组工作','2026-09-06','2026-09-06',[self.sid],[{'title':'任务','min_minutes':30,'max_minutes':60}],attachments=[f],group_id=gid)['collection_id']
        self.assertEqual(self.s.collections(self.other),[])
        self.assertEqual(len(self.s.collection_participants(self.normal,bid,group_id=gid)),1)
        with self.assertRaises(DomainError):self.s.collection_participants(self.normal,bid,group_id=other)
        with self.assertRaises(DomainError):self.s.create_collection(self.normal,bid,'越权','2026-09-06','2026-09-06',[self.sid],[{'title':'任务','min_minutes':30,'max_minutes':60}])
        row=self.s.collections(self.normal)[0]
        with self.assertRaises(DomainError):self.s.collection_attachment(self.other,cid,row['attachments'][0]['attachment_id'])
        sub=self.s.submit_hours(self.student,row['items'][0]['item_id'],'2026-09-06',60,attachments=[f])['submission_id']
        self.s.set_group_member(self.sup,gid,self.sid,False,'2026-09-06','调组')
        self.s.decide_hours(self.normal,sub,True,1);self.assertEqual(self.s.hours(self.student)['total_minutes'],0)
        self.s.decide_hours(self.sup,sub,True,2)
        self.assertEqual(self.s.hours(self.student,group_id=gid)['total_minutes'],60)
        self.assertEqual(self.s.hours(self.student,group_id=other)['total_minutes'],0)
        self.assertEqual(sum(r['total_minutes'] for r in self.s.hours(self.student)['group_totals']),60)
        term=self.s.save_semester(self.sup,'导出小组','2026-09-01',2)['semester_id']
        for group_filter in (None,gid,other):
            book=load_workbook(BytesIO(export_hours(self.s,self.sup,term,bid,group_filter)),data_only=True)
            summary=book['分队时长统计'];totals=[row[-1].value for row in summary if row[0].value=='全队合计']
            self.assertEqual(totals,[0 if group_filter==other else 1])
        self.s.set_collection_managers(self.sup,self.normal['id'],[],'撤权',[])
        with self.assertRaises(DomainError):self.s.submission_attachment(self.normal,sub,self.s.collections(self.student)[0]['submissions'][0]['attachments'][0]['attachment_id'])

    def test_group_scope_request_approval_and_rename_identity(self):
        gid=self.group();bid=self.b['干部']['business_id'];req=self.s.request_collection_scope(self.normal,[],'申请小组',[gid])['request_ids'][0]
        self.s.decide_collection_scope(self.sup,req,True,1)
        self.assertEqual(len(self.s.collection_management(self.normal)['group_grants']),1)
        migrate(self.settings);self.assertEqual(self.rows("SELECT business_id FROM business WHERE name='干部及干事'")[0]['business_id'],bid)
        self.assertFalse(self.rows("SELECT * FROM business WHERE name='干部'"))

    def test_roster_conflicts_replay_password_and_hours_unchanged(self):
        book=Workbook();sheet=book.active;sheet.title='综运';sheet.append(['姓名','学号','学院']);sheet.append(['张三',self.sid,'新学院']);sheet.append(['新队员','0012345678901','测试学院']);sheet.append(['缺失',None,None])
        data=BytesIO();book.save(data);encoded=base64.b64encode(data.getvalue()).decode()
        before=self.rows('SELECT * FROM student WHERE student_id=%s',(self.sid,))[0];minutes=self.s.hours(self.student)['total_minutes']
        preview=self.s.import_members(self.sup,encoded);self.assertEqual(preview['new_students'],1);self.assertEqual(len(preview['errors']),1)
        with self.assertRaises(DomainError):self.s.import_members(self.sup,encoded,commit=True,initial_password='12345678')
        result=self.s.import_members(self.sup,encoded,decisions={'综运:4':{'skip':True}},commit=True,initial_password='12345678')
        self.assertEqual(result['imported'],2);self.assertEqual(before,self.rows('SELECT * FROM student WHERE student_id=%s',(self.sid,))[0])
        self.assertTrue(self.s.import_members(self.sup,encoded,decisions={'综运:4':{'skip':True}},commit=True,initial_password='12345678')['replayed'])
        self.assertEqual(self.s.hours(self.student)['total_minutes'],minutes)
        from app.security import verify_password
        self.assertTrue(verify_password('12345678',self.rows("SELECT password_hash FROM student WHERE student_id='0012345678901'")[0]['password_hash']))

if __name__=='__main__':unittest.main()
